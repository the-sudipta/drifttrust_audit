from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
C=json.loads((ROOT/'qa/content.json').read_text(encoding='utf-8'))
activities={
'route':'Select a task. The highlighted code path changes between static presentation, live Worker processing and offline experiment. Reset to paper results; no network calls are made.',
'journey':'Choose example, replay or own CSV and toggle label availability. The recommended controls and learning eligibility update. Reset to first-time visitor; do not imply the offline guide contacts the live service.',
'split':'Assign sources for scaler, anchors and initial weights. Report protocol leakage or valid selection. Reset to an intentionally incorrect assignment; use actual fixed partition counts.',
'score':'Choose a real example and change packet rate. Compute the actual frozen MLP locally, show score/flag and edited-input caveat. Reset restores the original row. Zero and large rates are boundary probes, not new measurements.',
'detector':'Change recent/reference error counts, attribution difference, available labels and spacing. Show signal and candidate eligibility separately using the actual thresholds. Reset to the worked case; equality does not pass strict drift thresholds.',
'order':'Use selects to place prediction, feedback, training, checking and activation in order. Show the first order violation and why it matters. Reset scrambles the steps; no retroactive prediction rewriting is allowed.',
'gate':'Change ECI, loss/miss increases and policy mode. Show failed checks, activation/rejection and active model identity. Reset to the failed-loss example; equality passes gate limits.',
'audit':'Verify a copied real primary-run AJR using canonical JSON SHA-256; edit status and reverify. Show mismatch and explain trusted-head limits. Reset restores the untouched copy; no original records change.',
'metrics':'Adjust TP/FP/FN counts and compute precision, recall and F1. Show each denominator and distinguish F1 from accuracy. Reset to 80/20/40; empty denominators return zero by a stated teaching convention.'
}
brief=['# DriftTrust Audit code and live lab teaching brief','',
'## Purpose and learner',C['intro'],C['version'],
'Learner: Sudipta, CSE graduate and instructor. Prerequisites: basic functions, HTTP and elementary classification; every project-specific term is defined. Mastery: trace a request, operate all three input routes, distinguish scoring from learning, interpret governance records, reproduce the experiment and explain its limitations.',
'## Deliverables','This completed brief; a substantive Word dossier; an offline single-file learning_module.html. Store everything in this new topic directory. Preserve the deployed repository and supplied research artifacts.',
'## Language and depth','English technical explanations with natural বাংলা bridges. Nine ordered modules, five labelled teaching steps per module, a worked application, an activity beyond recall, a quiz, and an explain-back checklist. Simple arithmetic first; no unexplained symbolic notation. Actual checkpoint outputs are marked separately from teaching examples.',
'## Module specifications']
for i,m in enumerate(C['modules'],1):
    brief += [f"### {i} {m['title']}",f"Outcome: {m['lead']} Accent: {m['accent']}.",f"Question and prediction: {m['prediction']}",f"Action, consequence, explanation, reset and boundary: {activities[m['activity']]}",f"Explain-back: {m['explain']}",'Sources: '+', '.join(m['refs'])+'.']
brief+=['## Evidence and claim ledger',
'Primary sources are the inspected repository snapshot, actual model.json, examples.json, dataset.json, execution/results files and recorded hosted verification. Dataset attribution: B. S. Sharmila and Rohini Nagapadma, RT-IoT2022, UCI, DOI 10.24432/C5P338, CC BY 4.0; https://archive.ics.uci.edu/dataset/942/rt-iot2022. Figures and reported F1 come from the executed repository experiment. No new empirical result is inferred from an interactive teaching calculation.',
'Resolve contradictions: legacy label 1 is legitimate, current label 1 is attack; legacy temporal MLP is not an LSTM; public replay has 1,280 rows whereas full shifted test has 7,972. Recorded checks are historical evidence, not continuous monitoring. No claim of full ADWIN, SHAP, ISO certification, network enforcement, independent population trials or reviewer acceptance.',
'## Interaction and accessibility',
'One active lesson card, nine navigation items, distinct accents, local Nirmala UI Bengali font fallback, high contrast, keyboard controls, reduced-motion support, 27 shuffled review cards, glossary search and explicit note export. Retain answers and drafts in JavaScript memory only; reload clears them. Progress is self-assessed and never inferred from typing.',
'## Acceptance criteria',
'All three files contain substantive consistent content. Read the completed Word dossier back before authoring HTML. Render Word and inspect every page. Run the skill static checker. Exercise every activity, reset/boundary behavior, wrong/correct quiz, draft retention and safe angle brackets, notes export, shuffled/filtered/restored deck and reload reset in a browser with offline file loading. Inspect laptop/tablet/phone layouts and confirm no unexpected requests or console errors. Report precise verification limitations if a required check cannot run.']
(ROOT/'docs/DriftTrust_learning_module_instructions.md').write_text('\n\n'.join(brief),encoding='utf-8')

doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','List Number']:
    st=doc.styles[name];st.font.name='Nirmala UI';st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:cs'),'Nirmala UI')
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.14
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(26)
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(15)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
doc.core_properties.title=C['title'];doc.core_properties.subject='Source-grounded explanation and visitor walkthrough';doc.core_properties.author='DriftTrust Audit project learning guide'
header=sec.header.paragraphs[0];header.text='DriftTrust Audit     Code and visitor guide';header.style='Caption'
header.runs[0].font.color.rgb=RGBColor(0,0,0)
footer=sec.footer.paragraphs[0];footer.alignment=2
footer.add_run('Page ');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
def p(s):return doc.add_paragraph(s)
def h(s):doc.add_heading(s,2)
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for cell,text in zip(t.rows[0].cells,headers):cell.text=text
    for row in rows:
        for cell,text in zip(t.add_row().cells,row):cell.text=text
    for ri,row in enumerate(t.rows):
        for j,cell in enumerate(row.cells):
            cell.width=Inches(widths[j]);pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DFEAF2' if ri==0 else ('F5F7FA' if ri%2==0 else 'FFFFFF'));pr.append(shade)
            mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'95');e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar)
            for pa in cell.paragraphs:
                pa.paragraph_format.space_after=Pt(3)
                for run in pa.runs:run.font.size=Pt(9);run.bold=ri==0
        if ri==0:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
    p('')
doc.add_paragraph(C['title'],style='Title');p(C['intro']);p(C['version'])
p('Read the visitor walkthrough first. Then follow the nine modules from architecture to evidence. The companion fieldbook contains interactive exercises; the published lab loads from GitHub Pages and runs its model in your browser.')
doc.add_heading('Visitor walkthrough',1)
for step in C['quickstart']:doc.add_paragraph(step,style='List Number')
p('Live lab: https://the-sudipta.github.io/drifttrust_audit/lab.html')
for i,m in enumerate(C['modules'],1):
    doc.add_heading(str(i)+' '+m['title'],1);p(m['lead'])
    h('গল্প and motivation');p(m['story'])
    h('সহজ Analogy');p(m['analogy'])
    h('Formal definition and code flow')
    for s in m['formal']:p(s)
    h('Math and worked application')
    for step,s in enumerate(m['worked'],1):p(str(step)+'. '+s)
    h('Interactive check');p(m['prediction']);p(activities[m['activity']])
    p('Recall question: '+m['quiz']['q']);p('Answer: '+m['quiz']['options'][m['quiz']['answer']]+' '+m['quiz']['why'])
    p('Explain back: '+m['explain']);p('Self-check: '+' '.join(m['checklist']))
    p('Source files: '+', '.join(m['refs']))
doc.add_heading('Repository map',1);table(['File or directory','Responsibility'],C['filemap'],[2.7,4.3])
doc.add_heading('Browser operation map',1);table(['Route','Purpose'],C['api'],[2.3,4.7])
doc.add_heading('Every page control and display',1);table(['Control or text','Meaning and consequence'],C['ui_elements'],[2.4,4.6])
for key,title in [('features','All twelve input attributes'),('config','All configured model settings'),('audit','Every audit-record field'),('figures','How to read all nine research figures')]:
    doc.add_heading(title,1)
    for name,meaning in C['reference'][key]:
        h(name);p(meaning)

doc.add_heading('Developer commands and safe interpretation',1)
for s in ['pnpm install --frozen-lockfile → install the Node dependency lock. pnpm build:pages then pnpm dev:pages → local static lab at http://localhost:8788/drifttrust_audit/lab.html with browser IndexedDB state. pnpm dev retains the optional server API for separate API tests.',
'python research/prepare.py → retrieve/check data and fit initial models; node research/run.mjs → execute streaming comparisons; python research/report.py → regenerate metrics and figures. Install requirements-research.lock.txt in Python 3.12 first.',
'pnpm test → scientific engine tests. With the local server running, node tests/api-check.mjs http://localhost:8787 → API behavior checks. pnpm build → static assets and Worker bundle. Build does not itself publish.',
'node scripts/verify-audit.mjs downloaded-session-audit.json → recompute canonical hashes and verify chain links. A trusted saved head strengthens later substitution detection.',
'The original main.py is a separate synthetic workflow and clears files in its configured audit/results output directories at startup. Do not confuse running it with starting the live service.']:
    p(s)
doc.add_heading('Troubleshooting the visitor experience',1)
for s in ['Service unavailable: open GitHub Pages or run pnpm build:pages and pnpm dev:pages; double-clicking lab.html alone cannot load its browser module worker and JSON assets. The offline fieldbook is intentionally different and works without a server.',
'No candidates: ensure you used replay or a labelled stream, supplied enough labels and met spacing/trigger rules. One suspicious flow does not force retraining.',
'Policy appears unchanged: select a policy and then start a new session. A selection alone does not mutate the existing session.',
'Unexpected score: compare known truth separately, inspect training-range warnings and remember scores are uncalibrated. Edited example values no longer have a validated dataset label.',
'CSV rejected: check exact header names, 12 numeric features, non-negative finite values, 0/1/blank labels, maximum 256 rows and 2 MB browser file size. The optional older server API has a separate 400 KB JSON body limit; the browser lab sends no model inputs to it.',
'Session expired or conflict: create a new session after expiry; a 409 requires refreshing current revision. Export useful evidence before replacing the session. The last-request retry protection is not an unlimited request-history cache.']:
    p(s)
doc.add_heading('Glossary',1)
for term,meaning in C['glossary']:p(term+' — '+meaning)
doc.add_heading('Sources and claim boundaries',1)
p('Current repository and browser migration: https://github.com/the-sudipta/drifttrust_audit. The scientific engine remains the same as the original measured snapshot. Module source-file locators identify the implementation behind each explanation.')
p('Dataset: B. S. Sharmila and Rohini Nagapadma. RT-IoT2022. UCI Machine Learning Repository. DOI 10.24432/C5P338. CC BY 4.0. https://archive.ics.uci.edu/dataset/942/rt-iot2022')
p('Measured findings: assets/research/RESEARCH_REPORT.md, dataset.json, results.json and deployment_verification.json. This dossier adds explanation, not a new benchmark run. Actual initial-model example outputs were recomputed directly from the published checkpoint for this guide.')
p('Teaching scenarios and policy/metric arithmetic are explicitly illustrative. The embedded score activity uses real frozen weights; editing a feature creates a sensitivity probe. The offline guide never modifies a lab model session or submits measurements. Its audit activity edits only a copied record.')
doc.save(ROOT/'docs/DriftTrust_learning_dossier.docx')
read=Document(ROOT/'docs/DriftTrust_learning_dossier.docx')
text='\n'.join(p.text for p in read.paragraphs)
(ROOT/'qa/dossier_readback.txt').write_text(text,encoding='utf-8')
for m in C['modules']:
    assert m['title'] in text
    assert all(s in text for s in m['formal'])
assert 'ACCEPTED_WITH_FLAGS' in text and '7,972' in text
print(json.dumps({'modules':len(C['modules']),'paragraphs':len(read.paragraphs),'words':len(text.split()),'readback_complete':True}))
